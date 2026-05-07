"""Generate figures + final paper docx for the Gita embedding project."""
import json
import os
from pathlib import Path

import matplotlib.pyplot as plt
import numpy as np
import pandas as pd
from sklearn.cluster import KMeans

ROOT = Path("/Users/gurubazawada/Desktop/yoga class")
PROJ = ROOT / "gita_embeddings"
FIGS = ROOT / "paper_figures"
FIGS.mkdir(exist_ok=True)

plt.rcParams.update({"font.family": "DejaVu Serif", "font.size": 10})

# ---------- Load data ----------
emb = pd.read_parquet(PROJ / "data" / "embeddings_sim.parquet")
emb["vec"] = emb["embedding"].apply(np.asarray)
X = np.vstack(emb["vec"].values)
axes_df = pd.read_parquet(PROJ / "data" / "axis_scores.parquet")
chap_cent = pd.read_csv(PROJ / "reports" / "centroids_by_chapter.csv", index_col=0)

# ---------- Fig 1: Chapter heatmap ----------
fig, ax = plt.subplots(figsize=(6.5, 5.2))
M = chap_cent.values
im = ax.imshow(M, cmap="viridis", vmin=M[M < 0.999].min(), vmax=1.0)
ax.set_xticks(range(18)); ax.set_xticklabels(range(1, 19), fontsize=8)
ax.set_yticks(range(18)); ax.set_yticklabels(range(1, 19), fontsize=8)
ax.set_xlabel("Chapter"); ax.set_ylabel("Chapter")
ax.set_title("Chapter × Chapter Centroid Cosine Similarity")
plt.colorbar(im, ax=ax, fraction=0.046)
plt.tight_layout()
plt.savefig(FIGS / "fig1_chapter_heatmap.png", dpi=180)
plt.close()

# ---------- Fig 2: UMAP scatter colored by chapter ----------
try:
    import umap
    reducer = umap.UMAP(n_components=2, metric="cosine", random_state=42, n_neighbors=15, min_dist=0.1)
    coords = reducer.fit_transform(X)
except Exception as e:
    print("UMAP failed, using PCA:", e)
    from sklearn.decomposition import PCA
    coords = PCA(n_components=2).fit_transform(X)

fig, ax = plt.subplots(figsize=(7, 5.5))
chapters = emb["chapter"].values
sc = ax.scatter(coords[:, 0], coords[:, 1], c=chapters, cmap="tab20", s=12, alpha=0.85)
ax.set_title("UMAP of all 700 verses (colored by chapter)")
ax.set_xlabel("UMAP-1"); ax.set_ylabel("UMAP-2")
cb = plt.colorbar(sc, ax=ax, ticks=range(1, 19))
cb.set_label("Chapter")
plt.tight_layout()
plt.savefig(FIGS / "fig2_umap.png", dpi=180)
plt.close()

# ---------- Fig 3: Speaker radar ----------
axis_cols = [c for c in axes_df.columns if c != "verse_id"]
df = axes_df.merge(emb[["verse_id", "speaker"]], on="verse_id")
mean_by_speaker = df.groupby("speaker")[axis_cols].mean()
speakers_to_show = ["KRISHNA", "ARJUNA", "SANJAYA"]
mean_by_speaker = mean_by_speaker.loc[[s for s in speakers_to_show if s in mean_by_speaker.index]]

# Normalize each axis to [-1, 1] across speakers for visibility
norm = mean_by_speaker.copy()
for c in axis_cols:
    mx = norm[c].abs().max()
    if mx > 0:
        norm[c] = norm[c] / mx

angles = np.linspace(0, 2 * np.pi, len(axis_cols), endpoint=False).tolist()
angles += angles[:1]
fig, ax = plt.subplots(figsize=(7.5, 7), subplot_kw=dict(polar=True))
colors = {"KRISHNA": "#1f77b4", "ARJUNA": "#d62728", "SANJAYA": "#2ca02c"}
for sp, row in norm.iterrows():
    vals = row.tolist() + [row.tolist()[0]]
    ax.plot(angles, vals, label=sp, color=colors.get(sp), linewidth=2)
    ax.fill(angles, vals, alpha=0.10, color=colors.get(sp))
ax.set_xticks(angles[:-1])
ax.set_xticklabels([c.replace("-", "↔\n") for c in axis_cols], fontsize=7)
ax.set_yticklabels([])
ax.set_title("Speaker Voice Across the 17 Concept Axes\n(positive = toward second pole)", pad=20)
ax.legend(loc="upper right", bbox_to_anchor=(1.25, 1.1), fontsize=9)
plt.tight_layout()
plt.savefig(FIGS / "fig3_speaker_radar.png", dpi=180, bbox_inches="tight")
plt.close()

# ---------- Fig 4: Chapter arc on action-renunciation + sorrow-equanimity ----------
arc_df = axes_df.merge(emb[["verse_id", "chapter"]], on="verse_id")
arc_means = arc_df.groupby("chapter")[["action-renunciation", "sorrow-equanimity", "ego-surrender"]].mean()

fig, ax = plt.subplots(figsize=(7, 4.2))
ax.axhline(0, color="grey", linewidth=0.7)
for col, color, label in [
    ("action-renunciation", "#1f77b4", "action ↔ renunciation"),
    ("sorrow-equanimity", "#d62728", "sorrow ↔ equanimity"),
    ("ego-surrender", "#2ca02c", "ego ↔ surrender"),
]:
    ax.plot(arc_means.index, arc_means[col], marker="o", linewidth=2, color=color, label=label)
ax.set_xlabel("Chapter"); ax.set_ylabel("Mean axis score")
ax.set_xticks(range(1, 19))
ax.set_title("Chapter-by-chapter arcs across three key axes")
ax.legend(fontsize=9, loc="best")
plt.tight_layout()
plt.savefig(FIGS / "fig4_chapter_arcs.png", dpi=180)
plt.close()

# ---------- Fig 5: 'yoga' polysemy via k-means on the verses containing 'yoga' ----------
mask = emb["english"].str.contains(r"\byoga\b", case=False, regex=True)
yoga = emb[mask].reset_index(drop=True)
Xy = np.vstack(yoga["vec"].values)
k = 4
km = KMeans(n_clusters=k, random_state=42, n_init=10).fit(Xy)
yoga["cluster"] = km.labels_
counts = yoga["cluster"].value_counts().sort_index()
# pick exemplar = closest verse to each cluster centroid
exemplars = []
for cid in range(k):
    members = yoga[yoga["cluster"] == cid]
    if members.empty:
        exemplars.append("")
        continue
    Xc = np.vstack(members["vec"].values)
    center = km.cluster_centers_[cid]
    dists = np.linalg.norm(Xc - center, axis=1)
    ex = members.iloc[int(np.argmin(dists))]
    exemplars.append(f"{ex['verse_id']}")

fig, ax = plt.subplots(figsize=(6.5, 4))
bars = ax.bar(range(k), counts.values, color=["#1f77b4", "#ff7f0e", "#2ca02c", "#d62728"][:k])
ax.set_xticks(range(k))
ax.set_xticklabels([f"Cluster {i}\n({exemplars[i]})" for i in range(k)], fontsize=9)
for b, v in zip(bars, counts.values):
    ax.text(b.get_x() + b.get_width() / 2, v + 0.1, str(v), ha="center", fontsize=9)
ax.set_ylabel("Number of verses")
ax.set_title('Polysemy of "yoga": k-means clusters on verses containing the word')
plt.tight_layout()
plt.savefig(FIGS / "fig5_yoga_polysemy.png", dpi=180)
plt.close()

print("Figures done.")

# ---------- Build the docx ----------
from docx import Document
from docx.shared import Inches, Pt
from docx.enum.text import WD_ALIGN_PARAGRAPH

doc = Document()
styles = doc.styles
normal = styles["Normal"]
normal.font.name = "Times New Roman"
normal.font.size = Pt(11)

for sec in doc.sections:
    sec.top_margin = Inches(0.8)
    sec.bottom_margin = Inches(0.8)
    sec.left_margin = Inches(0.9)
    sec.right_margin = Inches(0.9)

def add_h(text, level=1):
    p = doc.add_heading(text, level=level)
    for run in p.runs:
        run.font.name = "Times New Roman"
    return p

def add_p(text, italic=False, align=None):
    p = doc.add_paragraph()
    if align is not None:
        p.alignment = align
    r = p.add_run(text)
    r.font.name = "Times New Roman"
    r.font.size = Pt(11)
    if italic:
        r.italic = True
    return p

def add_fig(path, caption, width=6.0):
    doc.add_picture(str(path), width=Inches(width))
    last = doc.paragraphs[-1]
    last.alignment = WD_ALIGN_PARAGRAPH.CENTER
    cap = doc.add_paragraph()
    cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = cap.add_run(caption)
    r.italic = True
    r.font.size = Pt(9)
    r.font.name = "Times New Roman"

# Title
title = doc.add_paragraph()
title.alignment = WD_ALIGN_PARAGRAPH.CENTER
tr = title.add_run("Reading the Bhagavad Gītā Through an Embedding Model")
tr.bold = True; tr.font.size = Pt(16); tr.font.name = "Times New Roman"
sub = doc.add_paragraph()
sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
sr = sub.add_run("Guru Bazawada  |  SAST 3650  |  Final Paper")
sr.font.size = Pt(11); sr.font.name = "Times New Roman"

# Introduction
add_h("1. Introduction", 1)
add_p(
    "The Bhagavad Gītā is a 700-verse philosophical dialogue tucked inside the Mahābhārata, in "
    "which the warrior Arjuna freezes on the edge of a battlefield and his charioteer Krishna "
    "spends eighteen chapters talking him back into action. People have read it as a war manual, "
    "a quietist tract, a devotional handbook, and a metaphysics textbook — sometimes all at once. "
    "I wanted to know what an embedding model would do with it. If you hand a modern language "
    "model every verse of the Gītā and ask only for vectors, with no labels and no theology, will "
    "the structure that scholars have spent centuries describing actually fall out of the numbers? "
    "This paper says yes — mostly — and shows what it looks like."
)
add_p(
    "I used Google's Gemini Embedding 2 model to embed all 700 verses of Eknath Easwaran's "
    "translation, projected those embeddings onto seventeen concept axes I designed by hand, ran "
    "UMAP for visualization, and clustered the verses around polysemous words like \"yoga\" and "
    "\"dharma\". The whole pipeline is roughly 1,400 API calls and a couple of dataframes; the "
    "interesting part is the readings it produces, which I summarize in Section 3."
)

# Method
add_h("2. Method", 1)
add_p(
    "The pipeline has four stages. First, I parsed the Easwaran ePub into a JSON corpus of one "
    "verse per row, tagged with chapter, verse number, and speaker (Krishna, Arjuna, Sanjaya, or "
    "Dhritarashtra). Second, every verse was embedded twice using gemini-embedding-2-preview, "
    "which produces self-normalized 3072-dimensional vectors. Gemini 2 swaps the older "
    "task_type parameter for in-text task prefixes, so each verse went through the model once "
    "with a \"sentence similarity\" prefix (used for axes, UMAP, and centroids) and once with a "
    "\"retrieval document\" prefix (used for semantic search). A SHA-256 cache keyed on "
    "model | dimension | formatted-text means re-runs are free."
)
add_p(
    "Third, I built seventeen concept axes. Each axis is a pair of pole phrases — for example "
    "\"action ↔ renunciation\", \"fear ↔ fearlessness\", \"ego ↔ surrender\". I embed both poles, "
    "subtract one from the other to get an axis vector, and project every verse onto that vector "
    "by cosine similarity. The result is a number between roughly −0.2 and +0.2 telling you how "
    "much each verse leans toward one pole or the other. The seventeen axes were chosen to cover "
    "the practical, psychological, and metaphysical themes the text repeatedly returns to (karma "
    "vs. sannyāsa, rāga vs. vairāgya, ahaṃkāra vs. śaraṇāgati, and so on)."
)
add_p(
    "Fourth, for visualization I ran UMAP (cosine metric, n_neighbors = 15) on the 700×3072 matrix "
    "to project to two dimensions, computed mean-centroid cosine similarity matrices grouped by "
    "chapter and by speaker, and ran k-means on the verse subsets that contain particular English "
    "words (\"yoga\", \"dharma\", \"karma\", etc.) to see whether the model recovers the classical "
    "polysemy of those terms. None of the analysis looks at Sanskrit — the embeddings only "
    "ever see Easwaran's English translation."
)

# Findings
add_h("3. Findings", 1)

add_h("3.1 Chapter 1 and Chapter 16 are structural outliers", 2)
add_p(
    "The cleanest result in the whole project is the chapter centroid heatmap (Figure 1). I "
    "averaged the embeddings of every verse in each chapter and computed the cosine similarity "
    "between every pair of chapter means. Almost the entire matrix is above 0.95 — chapters of "
    "the Gītā mostly sound like other chapters of the Gītā. The exceptions jump out immediately. "
    "Chapter 1, where Arjuna falls apart on the battlefield, has the lowest similarity to almost "
    "every other chapter. Chapter 16, the catalogue of the \"demonic\", is the second-lowest. "
    "Chapter 11, the cosmic vision, is the third. The book breaks its register exactly where you "
    "would expect: once to show the problem (Arjuna's collapse), once to show the metaphysical "
    "ground (the cosmic form), and once to show what living on the surface-ego looks like. The "
    "tightest pair, by contrast, is chapters 7 ↔ 9 (cosine ≈ 0.989), both being Krishna's "
    "discourses on the nature of the divine."
)
add_fig(FIGS / "fig1_chapter_heatmap.png",
        "Figure 1. Cosine similarity between chapter-mean embeddings. Chapters 1, 11, and 16 are visibly cooler than the rest of the matrix.")

add_h("3.2 UMAP recovers thematic neighborhoods", 2)
add_p(
    "Projecting all 700 verses to two dimensions with UMAP (Figure 2) gives a map where verses "
    "that the model considers semantically similar are close together. Coloring by chapter shows "
    "that the chapters do not occupy disjoint regions — the Gītā genuinely circles back to the "
    "same ideas — but neighboring chapters tend to puddle near each other. Chapter 1's verses "
    "form a small, almost separate island in the lower-left; chapter 11 forms another. Chapters "
    "5 and 6 sit near the center of mass of the figure, which lines up with the next finding."
)
add_fig(FIGS / "fig2_umap.png",
        "Figure 2. UMAP projection (cosine, n_neighbors = 15) of all 700 verses, colored by chapter. Chapter 1 forms a visible outlier cluster.")

add_h("3.3 The action / renunciation tension dissolves into a third option", 2)
add_p(
    "The most famous practical question in the Gītā is: should I act in the world or renounce it? "
    "On the action–renunciation axis, the top \"action\" verses (BG 3.8, 18.47) all say keep "
    "going. The top \"renunciation\" verses do not say stop; they redefine renunciation. BG 18.11 "
    "is the cleanest statement of the resolution: \"As long as one has a body, one cannot "
    "renounce action altogether. True renunciation is giving up all desire for personal reward.\" "
    "Run a semantic search for \"renunciation of the fruits of action\" and the top five hits are "
    "BG 18.11, 2.47, 18.12, 18.2, and 2.49 — all minor variations of the same teaching. The "
    "embeddings cluster these together because they are saying the same thing. The chapter-arc "
    "plot (Figure 3) shows this graphically: action-renunciation drifts steadily upward across "
    "the book, peaking in chapters 5, 6, and 18, where Krishna is restating the resolution."
)
add_fig(FIGS / "fig4_chapter_arcs.png",
        "Figure 3. Chapter-mean scores on three axes. Sorrow peaks in chapter 1, equanimity rises through chapters 5–6, and the ego/surrender axis tilts decisively toward surrender by chapters 9, 12, and 18.")

add_h("3.4 Speakers have measurably different voices", 2)
add_p(
    "The radar plot (Figure 4) shows the mean axis score for Krishna, Arjuna, and Sanjaya across "
    "all seventeen axes (each axis normalized so the largest absolute value is one, just for "
    "readability). Krishna leans noticeably toward equanimity, peace, spirit, contentment, faith, "
    "and fearlessness — the \"ideal\" pole on most axes. Arjuna leans the opposite way on most of "
    "them: more pain, more world, more sorrow, more doubt. Sanjaya, the narrator, sits in the "
    "middle, which is exactly right for a framing voice that is reporting the conversation rather "
    "than participating in it. None of this was supplied to the model — speaker identity was "
    "stored as a label but never used during embedding. The voice differences come straight out "
    "of the language each speaker actually uses."
)
add_fig(FIGS / "fig3_speaker_radar.png",
        "Figure 4. Mean axis scores per speaker across the seventeen concept axes (per-axis normalized). Krishna and Arjuna are nearly mirror images; Sanjaya sits between them.")

add_h('3.5 "Yoga" decomposes into four distinct senses', 2)
add_p(
    "I pulled out every verse containing the English word \"yoga\" (25 verses) and ran k-means "
    "with k = 4 on their embeddings. The clusters separate cleanly (Figure 5) and they line up "
    "with the four classical senses of the word: yoga as general practice or discipline, yoga as "
    "the path to Self-realization, yoga as a divine title (yogeśvara, \"master of yoga\", as in "
    "Krishna's epithets in chapter 11), and yoga as union with the divine in the bhakti sense "
    "(BG 12.2). The model finds these distinctions automatically, which is striking — \"yoga\" in "
    "the Gītā never means the physical postures it means in a modern studio, and the embedding "
    "respects that. The same procedure applied to \"dharma\" recovers four senses too "
    "(sanātana dharma, the \"field of dharma\", dharma's historical decline, and svadharma)."
)
add_fig(FIGS / "fig5_yoga_polysemy.png",
        'Figure 5. K-means clusters (k = 4) over the 25 verses containing the word "yoga". Each cluster\'s exemplar verse is the verse closest to the cluster centroid.')

add_h("4. Limitations", 1)
add_p(
    "The embeddings see only the English translation. Sanskrit wordplay — \"yoga\" from "
    "yuj-, to yoke, audible throughout the original — is invisible to the model. Easwaran also "
    "softens some of the rougher rhetorical moves in the text (BG 2.37 reads almost cynically in "
    "literal English but is gentled here), and the model inherits the softening. The other big "
    "miss is dramatic pacing: BG 11.32 (\"I am time, destroyer of worlds\") is narratively the "
    "shock-point of the whole book, but the embedding sees it as one more verse about the "
    "divine, which it also is. Numerical similarity is not the same as literary force."
)

add_h("5. Conclusion", 1)
add_p(
    "Almost every structural claim a careful reader would make about the Bhagavad Gītā shows up "
    "in the embedding analysis without being told to: chapter 1 is the diagnostic problem, "
    "chapters 5–6 are the calm center, chapter 16 is the foil, chapter 11 is its own thing, and "
    "the four \"yogas\" really are four senses of one word. The teaching the model converges on "
    "is the one the text is most insistent about — do your work, do it well, stop demanding it "
    "save you, and discover, in no longer demanding, that you were already what you were looking "
    "for. Seventeen axes, seven hundred verses, one claim."
)

out = ROOT / "Gita_Embedding_Paper.docx"
doc.save(out)
print("Wrote", out)
